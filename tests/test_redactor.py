"""Small checks for the rules most likely to break. Run: python -m pytest tests"""
import docx

from redactor import detectors, fakes
from redactor.pipeline import leak_check, redact_file


def types_found(text):
    return [(text[s:e], t) for s, e, t, _ in detectors.detect_all([text])[0]]


def test_luhn_rejects_order_numbers():
    assert types_found("Order #4111111111111112 dispatched.") == []
    assert ("4111 1111 1111 1111", "CREDIT_CARD") in types_found("Refund on card 4111 1111 1111 1111.")


def test_public_bodies_are_kept():
    assert types_found("As per SEBI rules, the shares are listed on BSE Limited.") == []


def test_dob_needs_context():
    assert types_found("The meeting was held on 14/03/1988.") == []
    assert ("14/03/1988", "DOB") in types_found("Date of birth: 14/03/1988")


def test_ticket_and_version_numbers_are_not_pii():
    assert types_found("Ticket 123-45-678 closed after upgrading to version 1.2.3.4.") == []


def test_aadhaar_checksum():
    assert types_found("Aadhaar 2943 6593 3462") == []                        # invalid check digit
    assert ("2943 6593 3461", "AADHAAR") in types_found("Aadhaar 2943 6593 3461")


def test_name_variants_share_one_fake():
    fakes.reset()
    assert fakes.fake_for("Kushal Subbayya Hegde", "PERSON") == fakes.fake_for("Kushal Hegde", "PERSON")
    assert fakes.fake_for("KUSHAL HEGDE", "PERSON").isupper()


def test_phone_keeps_country_code():
    fakes.reset()
    assert fakes.fake_for("+91 20 4505 3237", "PHONE").startswith("+91 ")


def test_docx_round_trip_keeps_structure_and_leaks_nothing(tmp_path):
    src, dst = tmp_path / "in.docx", tmp_path / "out.docx"
    doc = docx.Document()
    doc.add_paragraph("Contact Rashmi Kulkarni at rashmi.kulkarni@gmail.com or +91 98765 43210.")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Registered Office: 201, Tower 2, Baner, Pune – 411 045"
    doc.save(src)
    log = redact_file(str(src), str(dst))
    out = docx.Document(dst)
    assert len(out.paragraphs) == 1 and len(out.tables) == 1
    assert leak_check(str(dst), log) == []
    assert "rashmi" not in " ".join(p.text for p in out.paragraphs).lower()
