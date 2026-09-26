"""Evaluate the redactor. Usage: python -m eval.evaluate <original.docx> <redacted.docx>

1. Hand-labelled prospectus paragraphs: dev set (used while tuning) and test set (never tuned on)
2. Synthetic ticket log: PII types missing from the prospectus + traps (order / ticket / version numbers)
3. Injection: known fake PII inserted into real prospectus paragraphs (recall in real context)
4. Leak check, image check and formatting check on the redacted file

A prediction counts as correct (TP) if it has the same type as a gold label and the spans overlap.
Accuracy is token-level: share of words correctly marked PII / not PII."""
import io
import json
import random
import re
import sys
from collections import Counter, defaultdict

import docx
from PIL import Image
from faker import Faker

from redactor import detectors, docx_io


def gold_spans(text, labels):
    spans, start_from = [], Counter()
    for pii_type, value in labels:
        start = text.find(value, start_from[value])     # next occurrence if a value repeats
        assert start >= 0, f"label not in text: {value}"
        spans.append((start, start + len(value), pii_type))
        start_from[value] = start + 1
    return spans


def overlaps(a, b):
    return a[0] < b[1] and b[0] < a[1]


def score(texts, golds, preds):
    """Per-type TP/FP/FN, token counts and the list of errors."""
    stats, tokens, errors = defaultdict(Counter), Counter(), []
    for text, gold, pred in zip(texts, golds, preds):
        used = set()
        for g in gold:
            hit = next((i for i, p in enumerate(pred) if i not in used and p[2] == g[2] and overlaps(p, g)), None)
            if hit is None:
                stats[g[2]]["fn"] += 1
                errors.append(f"FN {g[2]}: {text[g[0]:g[1]]!r}")
            else:
                used.add(hit)
                stats[g[2]]["tp"] += 1
        for i, p in enumerate(pred):
            if i not in used:
                stats[p[2]]["fp"] += 1
                errors.append(f"FP {p[2]}: {text[p[0]:p[1]]!r}")
        for m in re.finditer(r"\S+", text):
            tok = m.span()
            tokens["correct"] += any(overlaps(tok, g) for g in gold) == any(overlaps(tok, p) for p in pred)
            tokens["total"] += 1
    return stats, tokens, errors


def table(title, stats, tokens):
    rows = [f"### {title}", "", "| Type | TP | FP | FN | Precision | Recall | F1 |", "|---|---|---|---|---|---|---|"]
    total = Counter()
    for t in sorted(stats) + ["ALL (micro)"]:
        c = total if t == "ALL (micro)" else stats[t]
        total.update({} if t == "ALL (micro)" else c)
        p = c["tp"] / (c["tp"] + c["fp"]) if c["tp"] + c["fp"] else 0
        r = c["tp"] / (c["tp"] + c["fn"]) if c["tp"] + c["fn"] else 0
        f = 2 * p * r / (p + r) if p + r else 0
        rows.append(f"| {t} | {c['tp']} | {c['fp']} | {c['fn']} | {p:.1%} | {r:.1%} | {f:.1%} |")
    return "\n".join(rows + ["", f"Token accuracy: {tokens['correct'] / tokens['total']:.1%}", ""])


def labelled_set(title, gold, preds_by_text):
    texts = [g["text"] for g in gold]
    stats, tokens, errors = score(texts, [gold_spans(g["text"], g["pii"]) for g in gold], [preds_by_text[t] for t in texts])
    return table(title, stats, tokens), errors


def injection(texts):
    """Append a sentence of known fake PII to 30 random paragraphs, run the full document, measure recall."""
    fake, rng = Faker("en_IN"), random.Random(7)
    Faker.seed(7)
    chosen, gold = rng.sample(range(len(texts)), 30), {}
    texts = list(texts)
    for i in chosen:
        values = {"PERSON": f"{fake.first_name()} {fake.last_name()}", "EMAIL": fake.email(),
                  "PHONE": "+91 " + fake.msisdn()[3:], "CREDIT_CARD": fake.credit_card_number(),
                  "IP_ADDRESS": fake.ipv4_public(), "PAN": fake.bothify("???P?####?").upper()}
        texts[i] += (f" Contact {values['PERSON']} at {values['EMAIL']} or {values['PHONE']}; card"
                     f" {values['CREDIT_CARD']} from IP {values['IP_ADDRESS']}, PAN {values['PAN']}.")
        gold[i] = [[t, v] for t, v in values.items()]
    preds = detectors.detect_all(texts)
    found = Counter()
    for i, labels in gold.items():
        for s, e, t in gold_spans(texts[i], labels):
            found[t, any(p[2] == t and overlaps(p, (s, e)) for p in preds[i])] += 1
    rows = ["### Injection test (recall of PII planted in 30 real paragraphs)", "", "| Type | Caught | Recall |", "|---|---|---|"]
    rows += [f"| {t} | {found[t, True]}/30 | {found[t, True] / 30:.0%} |" for t in sorted({t for t, _ in found})]
    return "\n".join(rows + [""])


def file_checks(original, redacted):
    """Leaks of values found by simple independent regexes, readable ID numbers in images, structure."""
    src, out = docx.Document(original), docx.Document(redacted)
    src_text = "\n".join(t for _, t in docx_io.paragraphs(src))
    out_text = "\n".join(t for _, t in docx_io.paragraphs(out))
    rows = ["### File checks", ""]
    for name, pattern in {"emails": r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+", "+91 phones": r"\+\s?91[\s\d-]{9,15}\d",
                          "SEBI reg. nos": r"\bIN[A-Z]\d{9}\b", "CINs": r"\b[LU]\d{5}[A-Z]{2}\d{4}[A-Z]{3}\d{6}\b"}.items():
        values = set(re.findall(pattern, src_text))
        rows.append(f"- Leak check, {name}: {sum(v in out_text for v in values)} of {len(values)} still present")
    images = [Image.open(io.BytesIO(p.blob)).convert("RGB") for p in out.part.package.parts if p.content_type.startswith("image/")]
    blank = sum(img.getextrema() == ((128, 128), (128, 128), (128, 128)) for img in images)
    rows.append(f"- Image check: {blank} of {len(images)} images are plain grey boxes (ID cards, logos, QR code removed)")
    rows.append(f"- Structure: paragraphs {len(src.paragraphs)} -> {len(out.paragraphs)}, tables {len(src.tables)} -> {len(out.tables)}")
    return "\n".join(rows + [""])


if __name__ == "__main__":
    original, redacted = sys.argv[1], sys.argv[2]
    texts = [t for _, t in docx_io.paragraphs(docx.Document(original))]
    preds = dict(zip(reversed(texts), reversed(detectors.detect_all(texts))))   # text -> spans (first occurrence)
    # ticket lines, and any gold paragraph read slightly differently from the document, are detected separately
    gold_texts = [g["text"] for f in ("dev", "test", "tickets") for g in json.load(open(f"eval/gold_{f}.json", encoding="utf-8"))]
    missing = [t for t in dict.fromkeys(gold_texts) if t not in preds]
    preds.update(zip(missing, detectors.detect_all(missing)))

    report, errors = [], []
    for title, path in [("Prospectus dev set (125 paragraphs, used while tuning)", "eval/gold_dev.json"),
                        ("Prospectus test set (75 paragraphs, never tuned on)", "eval/gold_test.json"),
                        ("Synthetic ticket log (18 tickets incl. traps)", "eval/gold_tickets.json")]:
        section, errs = labelled_set(title, json.load(open(path, encoding="utf-8")), preds)
        report.append(section)
        errors += [f"[{title.split(' (')[0]}] {e}" for e in errs]
    report += [injection(texts), file_checks(original, redacted), "### Errors", ""] + [f"- {e}" for e in errors]
    print("\n".join(report))
    open("eval/results.md", "w", encoding="utf-8").write("\n".join(report))
